"""知识库管理服务：解析上传文件（txt / pdf / word）+ 维护向量库索引。

数据流：
    管理员上传文件 → parse_upload 解析出纯文本 → 存 kb_document（数据源）
    → index_document 切分并写入 Chroma（检索用）

所以后台的「数据管理器」改的是 kb_document，向量库始终跟着它同步，
删除/修改文档时不会留下过期分片。
"""
import io
import os

from 智能体.admin import store
from 智能体.config_handler import chroma_conf
from 智能体.logger_hander import logger
from 智能体.path_toll import get_abs_path
from 智能体.rag.vector_store import VectorStoreService

# 首次启动若一个分类都没有，就用这个默认分类兜底
DEFAULT_CATEGORY = "扫地机器人"

# 允许上传的扩展名 → 来源类型（与 store.KB_SOURCE_TYPES 对得上）
ALLOWED_SUFFIXES = {".txt": "txt", ".pdf": "pdf", ".docx": "word"}

# 上传限制：单文件体积与 PDF 页数，避免超大文件拖垮解析与建索引
MAX_UPLOAD_MB = 10
MAX_UPLOAD_BYTES = MAX_UPLOAD_MB * 1024 * 1024
MAX_PDF_PAGES = 200


def _decode_text(data: bytes) -> str:
    """txt 先按 utf-8 解，失败再按 gbk（Windows 记事本常存成 gbk）"""
    for encoding in ("utf-8", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError(f"PDF 页数过多（{len(reader.pages)} 页），请拆分为 {MAX_PDF_PAGES} 页以内的文件后再上传")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx(data: bytes) -> str:
    from docx import Document as DocxDocument

    document = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    # 表格里的内容也要，否则商品参数表会整块丢失
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(part for part in parts if part.strip())


def parse_upload(filename: str, data: bytes) -> tuple[str, str]:
    """按扩展名解析上传文件，返回 (纯文本, 来源类型)；不支持的类型或解析失败抛 ValueError"""
    suffix = os.path.splitext(filename or "")[1].lower()
    source_type = ALLOWED_SUFFIXES.get(suffix)
    if source_type is None:
        raise ValueError(f"不支持的文件类型「{suffix or '未知'}」，仅支持 txt / pdf / docx")

    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError(
            f"文件过大（约 {len(data) / 1024 / 1024:.1f}MB），请压缩到 {MAX_UPLOAD_MB}MB 以内再上传"
        )

    try:
        if source_type == "txt":
            text = _decode_text(data)
        elif source_type == "pdf":
            text = _read_pdf(data)
        else:
            text = _read_docx(data)
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"文件解析失败，请确认文件未损坏、未加密（{e}）")

    text = (text or "").strip()
    if not text:
        raise ValueError("文件里没有解析到有效文本内容（可能是扫描件或空白文件）")
    return text, source_type


class KnowledgeBaseService:
    """把「数据库里的知识文档」与「向量库里的分片」当成一件事来维护"""

    def __init__(self):
        self.vector_store = VectorStoreService()

    def _index(self, doc_id: int, category: str, title: str, content: str) -> bool:
        """把一篇文档写入向量库，并把结果记进 kb_document.index_status（失败可重试）"""
        try:
            self.vector_store.index_document(doc_id, category, title, content)
            store.set_document_index_status(doc_id, store.KB_INDEX_DONE, None)
            return True
        except Exception as e:
            store.set_document_index_status(doc_id, store.KB_INDEX_FAILED, str(e))
            logger.error(f"[知识库]文档 {title}（id={doc_id}）建立索引失败：{str(e)}")
            return False

    def add_document(self, category: str, title: str, content: str, source_type: str) -> int:
        doc_id = store.create_kb_document(category, title, source_type, content)
        ok = self._index(doc_id, category, title, content)
        logger.info(f"[知识库]新增文档 {title}（分类 {category}，id={doc_id}），索引{'成功' if ok else '失败'}")
        return doc_id

    def update_document(self, doc_id: int, title: str, content: str) -> dict | None:
        """改标题或正文后重新索引，保证检索到的是最新内容"""
        row = store.update_kb_document(doc_id, title, content)
        if row:
            self._index(doc_id, row["category"], title, content)
            logger.info(f"[知识库]更新文档 {title}（id={doc_id}）并已重索引")
        return row

    def reindex_document(self, doc_id: int) -> dict | None:
        """对单篇文档重试建索引（失败后修复数据源，再点重试即可）"""
        row = store.get_kb_document(doc_id)
        if not row:
            return None
        self._index(doc_id, row["category"], row["title"], row["content"])
        return store.get_kb_document(doc_id)

    def delete_document(self, doc_id: int) -> dict | None:
        row = store.delete_kb_document(doc_id)
        if row:
            self.vector_store.delete_document(doc_id)
            logger.info(f"[知识库]删除文档 {row['title']}（id={doc_id}）并已清理分片")
        return row

    def delete_category(self, name: str) -> int:
        """连分类带文档一起删，并清掉这些文档的向量分片"""
        ids = store.delete_kb_category(name)
        for doc_id in ids:
            self.vector_store.delete_document(doc_id)
        logger.info(f"[知识库]删除分类 {name}，同时清理 {len(ids)} 篇文档")
        return len(ids)

    def rebuild_index(self) -> int:
        """按 kb_document 全量重建向量库（清空集合后逐篇重新切分写入），并逐篇更新索引状态"""
        self.vector_store.reset_index()
        count = 0
        for row in store.list_kb_documents(with_content=True):
            if self._index(row["id"], row["category"], row["title"], row["content"]):
                count += 1
        logger.info(f"[知识库]向量库已重建，成功 {count} 篇文档")
        return count


def ensure_kb_seeded() -> None:
    """首次启动：建一个默认商品分类，并把 智能体/data 里已有的 txt 知识导入管理器。

    只写 kb_document，不动向量库——向量库由 rebuild_index 统一重建。
    """
    if not store.list_kb_categories():
        store.create_kb_category(DEFAULT_CATEGORY)

    if store.kb_document_total() > 0:
        return

    data_dir = get_abs_path(chroma_conf["data_path"])
    if not os.path.isdir(data_dir):
        return

    for name in sorted(os.listdir(data_dir)):
        if not name.endswith(".txt"):
            continue
        with open(os.path.join(data_dir, name), "rb") as f:
            data = f.read()
        text, source_type = parse_upload(name, data)
        doc_id = store.create_kb_document(
            DEFAULT_CATEGORY, os.path.splitext(name)[0], source_type, text
        )
        logger.info(f"[知识库]导入已有知识文件 {name}（id={doc_id}）")


if __name__ == "__main__":
    # 命令行单独跑时，先确保表结构存在（Web 应用里由 lifespan 负责）
    store.init_db()
    ensure_kb_seeded()
    service = KnowledgeBaseService()
    print(f"分类：{[c['name'] + ' ' + str(c['doc_count']) for c in store.list_kb_categories()]}")
    print(f"重建索引：{service.rebuild_index()} 篇文档")
